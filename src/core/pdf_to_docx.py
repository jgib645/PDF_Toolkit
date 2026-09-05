import os
from pathlib import Path
from .converter import Converter, ProgressCallback
from ..utils.pdf_reader import PDFPageLoader
from ..utils.ocr import extract_text_from_image
from ..utils.logging import get_logger
from pdf2docx import Converter as Pdf2DocxConverter
from docx import Document


class PDFToDOCX(Converter):
    def convert(self, input_path: str, output_path: str, progress_cb: ProgressCallback = None) -> dict:
        logger = get_logger()
        loader = PDFPageLoader(input_path)
        total = loader.total_pages()
        summary = {'total': total, 'failed_pages': []}

        # Attempt full conversion first (handles all pages with a text layer)
        docx_converter = Pdf2DocxConverter(input_path)
        try:
            docx_converter.convert(output_path)
        except Exception as e:
            logger.error(f"pdf2docx failed: {e}")
        finally:
            docx_converter.close()

        # Open the resulting document once, so OCR fallback pages can be appended to it
        if os.path.exists(output_path):
            doc = Document(output_path)
        else:
            doc = Document()

        doc_modified = False

        for i in range(total):
            if progress_cb:
                progress_cb(i + 1, total)

            page_has_text = loader.has_text_layer(i)
            if not page_has_text:
                try:
                    img = loader.page_image(i)
                    text = extract_text_from_image(img)
                    doc.add_paragraph(text)
                    doc_modified = True
                except Exception as e:
                    logger.warning(f"OCR failed page {i + 1}: {e}")
                    summary['failed_pages'].append(i + 1)

        if doc_modified:
            doc.save(output_path)

        return summary